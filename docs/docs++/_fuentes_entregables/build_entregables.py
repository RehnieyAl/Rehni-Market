"""
Genera los 11 documentos DOCX de entregables SENA de RehniMarket a partir de las
fuentes Markdown en docs/_fuentes_entregables/.

Formato de TODOS los documentos (exigido por el enunciado):
  - Carta (letter), margenes 2.54 cm
  - Times New Roman, texto negro, encabezados negros
  - Interlineado 2.0 en el cuerpo, texto JUSTIFICADO
  - Tablas reales de Word: fondo blanco, bordes negros, encabezado en negrita, SIN sombreado
  - SIN colores, SIN emojis, SIN iconos, SIN fondos
  - Numeracion de paginas (pie, centrada); la portada NO lleva numero
  - Tabla de contenido actualizable (campo TOC de Word)
  - Lista de tablas (siempre que haya tablas); lista de figuras SOLO si hay figuras reales

Convenciones de las fuentes .md (controladas por el autor):
  # PORTADA ... <!-- PAGEBREAK -->        -> portada centrada, sin numero de pagina
  <!-- TOC -->                            -> campo TOC de Word
  <!-- LISTA-TABLAS -->                   -> lista de tablas autogenerada
  <!-- LISTA-FIGURAS -->                  -> lista de figuras autogenerada (solo si hay)
  <!-- PAGEBREAK -->                      -> salto de pagina
  # / ## / ###                            -> encabezados (Heading 1/2/3)
  [[TABLA]] Titulo de la tabla.           -> seguido de una tabla Markdown (| ... |)
  [[FIGURA]] Titulo de la figura.         -> seguido de un bloque ``` ... ``` (ASCII)
  ``` ... ```                             -> bloque de codigo monoespaciado
  - / 1.                                  -> listas
  parrafo normal                          -> parrafo justificado con sangria de primera linea
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "/home/rehnieyal/Escritorio/Rehni-Market"
SRC = os.path.join(ROOT, "docs", "_fuentes_entregables")

TARGETS = {
    "01_Plataforma_e_Infraestructura":  "01_MODULO_A_DESPLIEGUE/01_Plataforma_e_Infraestructura.docx",
    "02_Plan_Migracion_y_Backups":      "01_MODULO_A_DESPLIEGUE/02_Plan_Migracion_y_Backups.docx",
    "03_Despliegue_y_Publicacion":      "01_MODULO_A_DESPLIEGUE/03_Despliegue_y_Publicacion.docx",
    "04_Usuarios_Roles_y_Permisos":     "01_MODULO_A_DESPLIEGUE/04_Usuarios_Roles_y_Permisos.docx",
    "05_Documentacion_Tecnica_y_Manuales": "01_MODULO_A_DESPLIEGUE/05_Documentacion_Tecnica_y_Manuales.docx",
    "06_Pruebas_Aceptacion_y_Entrega":  "01_MODULO_A_DESPLIEGUE/06_Pruebas_Aceptacion_y_Entrega.docx",
    "07_Calidad_ISO_PSP":               "02_MODULO_B_CALIDAD/07_Calidad_ISO_PSP.docx",
    "08_Requisitos_No_Funcionales":     "02_MODULO_B_CALIDAD/08_Requisitos_No_Funcionales.docx",
    "09_Informe_Evaluacion_Calidad":    "02_MODULO_B_CALIDAD/09_Informe_Evaluacion_Calidad.docx",
    "10_Plan_Mejora_Continua":          "02_MODULO_B_CALIDAD/10_Plan_Mejora_Continua.docx",
    "11_Informe_Avance_90":             "03_AVANCE_PROYECTO/11_Informe_Avance_90.docx",
}

BLACK = RGBColor(0, 0, 0)
FONT = "Times New Roman"
MONO = "Courier New"
INLINE_RE = re.compile(r'(\*\*.+?\*\*|`.+?`)')


def _style_run(run, size=12, bold=False, mono=False, italic=False):
    name = MONO if mono else FONT
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name); rFonts.set(qn('w:hAnsi'), name); rFonts.set(qn('w:cs'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK


def _page_number_field(paragraph):
    run = paragraph.add_run()
    for t, txt in (('begin', None), ('instr', 'PAGE'), ('end', None)):
        if t == 'instr':
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        else:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), t)
        run._r.append(el)
    _style_run(run, size=10)


def _toc_field(paragraph):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = r'TOC \o "1-3" \h \z \u'
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = "Actualice este campo: clic derecho sobre esta linea > Actualizar campos > Actualizar toda la tabla."
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for el in (b, i, s, t, e):
        run._r.append(el)
    _style_run(run, size=11)


def _table_borders_black(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def setup_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0); pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for lvl, sz in ((1, 14), (2, 13), (3, 12)):
        st = doc.styles[f'Heading {lvl}']
        st.font.name = FONT; st.font.size = Pt(sz)
        st.font.bold = True; st.font.italic = False
        st.font.color.rgb = BLACK
        p = st.paragraph_format
        p.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        p.space_before = Pt(12 if lvl == 1 else 8); p.space_after = Pt(6)
        p.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def setup_section(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.59); sec.page_height = Cm(27.94)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    _page_number_field(fp)


def add_paragraph(doc, text, indent=True, bullet=False, justify=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    if bullet:
        pf.left_indent = Cm(0.75); pf.first_line_indent = Cm(-0.35)
        text = "• " + text
    elif indent:
        pf.first_line_indent = Cm(1.27)
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            _style_run(p.add_run(part[2:-2]), bold=True)
        elif part.startswith('`') and part.endswith('`'):
            _style_run(p.add_run(part[1:-1]), mono=True, size=10)
        else:
            _style_run(p.add_run(part))
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.left_indent = Cm(0.5)
    pf.space_before = Pt(4); pf.space_after = Pt(4)
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        _style_run(p.add_run(ln.replace('\t', '    ')), mono=True, size=9)


def add_caption(doc, kind, num, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(6); pf.space_after = Pt(3)
    pf.keep_with_next = True
    _style_run(p.add_run(f"{kind} {num}. "), bold=True, size=10)
    _style_run(p.add_run(text), size=10)


def add_md_table(doc, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.autofit = True
    _table_borders_black(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        pa = c.paragraphs[0]
        pa.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        _style_run(pa.add_run(h), bold=True, size=9)
    for row in rows:
        cells = t.add_row().cells
        for i in range(len(header)):
            val = row[i] if i < len(row) else ""
            c = cells[i]
            c.text = ""
            pa = c.paragraphs[0]
            pa.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            for part in INLINE_RE.split(val):
                if not part:
                    continue
                if part.startswith('**') and part.endswith('**'):
                    _style_run(pa.add_run(part[2:-2]), bold=True, size=9)
                elif part.startswith('`') and part.endswith('`'):
                    _style_run(pa.add_run(part[1:-1]), mono=True, size=8)
                else:
                    _style_run(pa.add_run(part), size=9)
    hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true')
    t.rows[0]._tr.get_or_add_trPr().append(hdr)


def _split_md_row(line):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    return [c.strip() for c in line.split('|')]


def prescan(lines):
    """Devuelve (tablas, figuras) = listas de (numero, caption) en orden."""
    tablas, figuras = [], []
    tc = fc = 0
    for idx, ln in enumerate(lines):
        s = ln.strip()
        m = re.match(r'^\[\[TABLA\]\]\s*(.+)$', s)
        if m:
            tc += 1; tablas.append((tc, m.group(1).strip()))
        m = re.match(r'^\[\[FIGURA\]\]\s*(.+)$', s)
        if m:
            fc += 1; figuras.append((fc, m.group(1).strip()))
    return tablas, figuras


def render(md_path, out_path):
    with open(md_path, encoding='utf-8') as fh:
        lines = fh.read().split('\n')

    tablas, figuras = prescan(lines)
    tnum = {cap: n for n, cap in tablas}
    fnum = {cap: n for n, cap in figuras}

    doc = Document()
    setup_styles(doc)
    setup_section(doc)

    portada = False
    i, n = 0, len(lines)
    ti = fi = 0

    while i < n:
        raw = lines[i]
        s = raw.strip()

        if s == '<!-- PAGEBREAK -->':
            doc.add_page_break(); i += 1; continue
        if s == '<!-- TOC -->':
            _toc_field(doc.add_paragraph()); i += 1; continue
        if s == '<!-- LISTA-TABLAS -->':
            if not tablas:
                add_paragraph(doc, "Este documento no contiene tablas.", indent=False)
            for num, cap in tablas:
                add_paragraph(doc, f"Tabla {num}. {cap}", indent=False, justify=False)
            i += 1; continue
        if s == '<!-- LISTA-FIGURAS -->':
            for num, cap in figuras:
                add_paragraph(doc, f"Figura {num}. {cap}", indent=False, justify=False)
            i += 1; continue
        if s.startswith('<!--') and '-->' not in s:
            i += 1
            while i < n and '-->' not in lines[i]:
                i += 1
            i += 1; continue
        if s.startswith('<!--'):
            i += 1; continue

        if s.startswith('```'):
            block = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                block.append(lines[i]); i += 1
            i += 1
            add_code_block(doc, block)
            continue

        if s.startswith('|') and i + 1 < n and re.match(r'^\|?\s*:?-{2,}', lines[i + 1].strip()):
            header = _split_md_row(lines[i])
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(_split_md_row(lines[i])); i += 1
            add_md_table(doc, header, rows)
            continue

        m = re.match(r'^\[\[TABLA\]\]\s*(.+)$', s)
        if m:
            cap = m.group(1).strip()
            add_caption(doc, "Tabla", tnum.get(cap, '?'), cap)
            i += 1; continue

        m = re.match(r'^\[\[FIGURA\]\]\s*(.+)$', s)
        if m:
            cap = m.group(1).strip()
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            block = []
            if j < n and lines[j].strip().startswith('```'):
                j += 1
                while j < n and not lines[j].strip().startswith('```'):
                    block.append(lines[j]); j += 1
                j += 1
            add_code_block(doc, block)
            add_caption(doc, "Figura", fnum.get(cap, '?'), cap)
            i = j
            continue

        mh = re.match(r'^(#{1,4})\s+(.*)$', s)
        if mh:
            lvl = len(mh.group(1)); text = mh.group(2).strip()
            if text.upper() == 'PORTADA':
                portada = True; i += 1; continue
            if portada and lvl >= 2:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
                _style_run(p.add_run(text.replace('**', '')), bold=True, size=16)
                i += 1; continue
            portada = False
            h = doc.add_heading(level=min(lvl, 3))
            h.text = ""
            _style_run(h.add_run(text), bold=True, size=14 if lvl == 1 else (13 if lvl == 2 else 12))
            h.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1; continue

        if s in ('---', '***', '___'):
            if portada:
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1; continue

        if not s:
            i += 1; continue

        if portada:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            big = s.startswith('##')
            clean = s.lstrip('#').strip().replace('**', '')
            _style_run(p.add_run(clean), bold=(clean.isupper() or big),
                       size=16 if big else 12)
            i += 1; continue

        if s.startswith('- '):
            add_paragraph(doc, s[2:].strip(), bullet=True); i += 1; continue
        if re.match(r'^\d+\.\s+', s):
            add_paragraph(doc, s, bullet=True); i += 1; continue

        add_paragraph(doc, s)
        i += 1

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    return len(tablas), len(figuras)


def main():
    only = sys.argv[1:] or list(TARGETS)
    for key in only:
        md = os.path.join(SRC, key + ".md")
        if not os.path.exists(md):
            print("  ! falta fuente:", md); continue
        out = os.path.join(ROOT, "docs", TARGETS[key])
        nt, nf = render(md, out)
        print(f"OK  {TARGETS[key]:<62}  tablas={nt:<3} figuras={nf}")


if __name__ == "__main__":
    main()
